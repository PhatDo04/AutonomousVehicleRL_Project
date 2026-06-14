# -*- coding: utf-8 -*-
"""Re-extract the CURRENT DoAnTotNghiep_PhatDT.docx -> full text + tables + image count.
Walks body in document order so tables appear where they actually sit."""
import sys
from docx import Document
from docx.oxml.ns import qn

PATH = "DoAnTotNghiep_PhatDT.docx"
doc = Document(PATH)

out = []
para_idx = 0
tbl_idx = 0

body = doc.element.body
# map xml element -> python object
para_map = {p._p: p for p in doc.paragraphs}
tbl_map = {t._tbl: t for t in doc.tables}

out.append("############ PART A+B: BODY IN DOCUMENT ORDER (style | text) ############")
img_total = 0

from docx.text.paragraph import Paragraph
from docx.table import Table

def emit_para(child):
    global para_idx, img_total
    p = para_map.get(child)
    if p is None:
        p = Paragraph(child, doc)  # SDT-nested para not in top-level map
    para_idx += 1
    style = p.style.name if p.style else "?"
    txt = p.text
    blips = child.findall('.//' + qn('a:blip'))
    nimg = len(blips)
    img_total += nimg
    tag = f"[P{para_idx:03d}|{style}]"
    if nimg:
        tag += f"[IMG x{nimg}]"
    if txt.strip() or nimg:
        out.append(f"{tag} {txt}")

def emit_table(child):
    global tbl_idx, img_total
    t = tbl_map.get(child)
    if t is None:
        t = Table(child, doc)
    tbl_idx += 1
    out.append(f"\n===== TABLE T{tbl_idx:02d} ({len(t.rows)} rows x {len(t.columns)} cols) =====")
    for ri, row in enumerate(t.rows):
        cells = []
        seen = set()
        for c in row.cells:
            cid = id(c._tc)
            if cid in seen:
                continue
            seen.add(cid)
            cells.append(" ".join(c.text.split()))
        out.append(f"  R{ri:02d}: " + " | ".join(cells))
    blips = child.findall('.//' + qn('a:blip'))
    if blips:
        img_total += len(blips)
        out.append(f"  [TABLE IMG x{len(blips)}]")
    out.append("===== /TABLE =====\n")

def walk(parent):
    """Walk children in document order; descend into SDT content controls."""
    for child in parent.iterchildren():
        if child.tag == qn('w:p'):
            emit_para(child)
        elif child.tag == qn('w:tbl'):
            emit_table(child)
        elif child.tag == qn('w:sdt'):
            content = child.find(qn('w:sdtContent'))
            if content is not None:
                out.append(f"[--- SDT content control ---]")
                walk(content)
                out.append(f"[--- /SDT ---]")

walk(body)

out.append(f"\n############ SUMMARY ############")
out.append(f"paragraphs={para_idx}  tables={tbl_idx}  inline_images(blips)={img_total}")

text = "\n".join(out)
with open("_docx_current.txt", "w", encoding="utf-8") as f:
    f.write(text)
print(f"wrote _docx_current.txt  paragraphs={para_idx} tables={tbl_idx} images={img_total}")
